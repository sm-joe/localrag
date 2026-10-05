from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedDocument:
    filename: str
    content_type: str
    text: str


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
}


def parse_document(
    file_path: str | Path,
    content_type: str,
) -> ParsedDocument:
    path = Path(file_path)

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {path.suffix}"
        )

    extension = path.suffix.lower()

    if extension == ".pdf":
        text = _parse_pdf(path)
    elif extension == ".docx":
        text = _parse_docx(path)
    elif extension in {".txt", ".md"}:
        text = path.read_text(
            encoding="utf-8",
            errors="replace",
        )
    else:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    normalized_text = _normalize_text(text)

    if not normalized_text:
        raise ValueError(
            "The document contains no extractable text."
        )

    return ParsedDocument(
        filename=path.name,
        content_type=content_type,
        text=normalized_text,
    )


def _parse_pdf(path: Path) -> str:
    import fitz

    pages: list[str] = []

    with fitz.open(path) as document:
        for page in document:
            pages.append(page.get_text())

    return "\n\n".join(pages)


def _parse_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n\n".join(paragraphs)


def _normalize_text(text: str) -> str:
    lines = [
        line.strip()
        for line in text.splitlines()
    ]

    return "\n".join(
        line
        for line in lines
        if line
    ).strip()