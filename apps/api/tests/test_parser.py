from pathlib import Path

import pytest

from app.services.parser import (
    MAX_DOCX_PARAGRAPHS,
    MAX_EXTRACTED_TEXT_CHARS,
    MAX_PDF_PAGES,
    parse_document,
)


def test_parse_txt_document(tmp_path: Path):
    file_path = tmp_path / "document.txt"

    file_path.write_text(
        "LocalRAG is a secure RAG application.",
        encoding="utf-8",
    )

    result = parse_document(
        file_path=file_path,
        content_type="text/plain",
    )

    assert result.filename == "document.txt"
    assert result.content_type == "text/plain"
    assert result.text == (
        "LocalRAG is a secure RAG application."
    )


def test_parse_md_document(tmp_path: Path):
    file_path = tmp_path / "document.md"

    file_path.write_text(
        "# LocalRAG\n\nRAG uses document retrieval.",
        encoding="utf-8",
    )

    result = parse_document(
        file_path=file_path,
        content_type="text/markdown",
    )

    assert result.filename == "document.md"
    assert result.content_type == "text/markdown"
    assert result.text == (
        "# LocalRAG\nRAG uses document retrieval."
    )


def test_parse_rejects_unsupported_extension(
    tmp_path: Path,
):
    file_path = tmp_path / "malicious.exe"

    file_path.write_bytes(
        b"not a supported document"
    )

    with pytest.raises(
        ValueError,
        match="Unsupported file type",
    ):
        parse_document(
            file_path=file_path,
            content_type="application/octet-stream",
        )


def test_parse_rejects_missing_file(
    tmp_path: Path,
):
    file_path = tmp_path / "missing.txt"

    with pytest.raises(
        ValueError,
        match="Document file does not exist",
    ):
        parse_document(
            file_path=file_path,
            content_type="text/plain",
        )


def test_parse_rejects_empty_text_file(
    tmp_path: Path,
):
    file_path = tmp_path / "empty.txt"

    file_path.write_text(
        "",
        encoding="utf-8",
    )

    with pytest.raises(
        ValueError,
        match="no extractable text",
    ):
        parse_document(
            file_path=file_path,
            content_type="text/plain",
        )


def test_parse_normalizes_whitespace(
    tmp_path: Path,
):
    file_path = tmp_path / "whitespace.txt"

    file_path.write_text(
        "\n\n  LocalRAG  \n\n\n  Qdrant  \n\n",
        encoding="utf-8",
    )

    result = parse_document(
        file_path=file_path,
        content_type="text/plain",
    )

    assert result.text == "LocalRAG\nQdrant"


def test_parse_rejects_excessive_extracted_text(
    tmp_path: Path,
):
    file_path = tmp_path / "large.txt"

    file_path.write_text(
        "A" * (MAX_EXTRACTED_TEXT_CHARS + 1),
        encoding="utf-8",
    )

    with pytest.raises(
        ValueError,
        match="too much extracted text",
    ):
        parse_document(
            file_path=file_path,
            content_type="text/plain",
        )


def test_parse_pdf_page_limit(
    tmp_path: Path,
    monkeypatch,
):
    import fitz

    file_path = tmp_path / "too-many-pages.pdf"

    document = fitz.open()

    for _ in range(MAX_PDF_PAGES + 1):
        document.new_page()

    document.save(file_path)
    document.close()

    with pytest.raises(
        ValueError,
        match="too many pages",
    ):
        parse_document(
            file_path=file_path,
            content_type="application/pdf",
        )


def test_parse_pdf_document(
    tmp_path: Path,
):
    import fitz

    file_path = tmp_path / "document.pdf"

    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "LocalRAG PDF security test.",
    )
    document.save(file_path)
    document.close()

    result = parse_document(
        file_path=file_path,
        content_type="application/pdf",
    )

    assert result.filename == "document.pdf"
    assert "LocalRAG PDF security test." in result.text


def test_parse_malformed_pdf(
    tmp_path: Path,
):
    file_path = tmp_path / "broken.pdf"

    file_path.write_bytes(
        b"This is not a valid PDF."
    )

    with pytest.raises(
        ValueError,
        match="PDF document could not be parsed",
    ):
        parse_document(
            file_path=file_path,
            content_type="application/pdf",
        )


def test_parse_docx_document(
    tmp_path: Path,
):
    from docx import Document

    file_path = tmp_path / "document.docx"

    document = Document()
    document.add_paragraph(
        "LocalRAG DOCX security test."
    )
    document.save(file_path)

    result = parse_document(
        file_path=file_path,
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        ),
    )

    assert result.filename == "document.docx"
    assert (
        "LocalRAG DOCX security test."
        in result.text
    )


def test_parse_docx_paragraph_limit(
    tmp_path: Path,
):
    from docx import Document

    file_path = tmp_path / "too-many-paragraphs.docx"

    document = Document()

    for index in range(
        MAX_DOCX_PARAGRAPHS + 1
    ):
        document.add_paragraph(
            f"Paragraph {index}"
        )

    document.save(file_path)

    with pytest.raises(
        ValueError,
        match="too many paragraphs",
    ):
        parse_document(
            file_path=file_path,
            content_type=(
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),
        )


def test_parse_malformed_docx(
    tmp_path: Path,
):
    file_path = tmp_path / "broken.docx"

    file_path.write_bytes(
        b"This is not a valid DOCX file."
    )

    with pytest.raises(
        ValueError,
        match="DOCX document could not be parsed",
    ):
        parse_document(
            file_path=file_path,
            content_type=(
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),
        )